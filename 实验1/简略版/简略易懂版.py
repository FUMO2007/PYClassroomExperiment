# -*- coding: utf-8 -*-
import os
import shutil
import openpyxl

folder = "C:/Users/34542/Desktop/实验课作业/简略版/"
NOT_ADMITTED = "未被录取"

src = folder + "原始报名表.xlsx"   # 输入：全部报名名单
res = folder + "报名结果.xlsx"     # 输入：录取名单
out = folder + "录取结果.xlsx"     # 输出：填好社团的名单

admitted = {}
ws = openpyxl.load_workbook(res).active
for name, major, phone, club in ws.iter_rows(min_row=2, values_only=True):
    admitted[(name, major, phone)] = club

# 以报名表为模板另存为输出文件，再去填「社团」列（不要动输入文件）
shutil.copy(src, out)

wb = openpyxl.load_workbook(out)
ws = wb.active
for row in ws.iter_rows(min_row = 2):
    name, major, phone, club = (c.value for c in row)
    row[3].value = admitted.get((name, major, phone), NOT_ADMITTED)
wb.save(out)
print("OK")