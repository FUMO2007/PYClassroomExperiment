# -*- coding: utf-8 -*-
"""实验1：合并文件夹中的两个 xlsx。

「原始报名表.xlsx」是全部报名名单，「报名结果.xlsx」是录取名单：
录取的学生在「社团」列显示录取社团，未录取的显示「未被录取」，结果写入「合并结果.xlsx」。
"""

import os

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

DIR = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(DIR, "原始报名表.xlsx")
RES = os.path.join(DIR, "报名结果.xlsx")
OUT = os.path.join(DIR, "合并结果.xlsx")
KEY, CLUB, NO = "姓名", "社团", "未被录取"


def read(path):
    """按文本读取（防止电话号码变成科学计数法），并去掉首尾空格。"""
    return pd.read_excel(path, dtype=str).fillna("").apply(lambda col: col.str.strip())


def merge(src, res):
    """以原始报名表为准合并，返回带最终「社团」列的结果。"""
    res = res[[KEY, CLUB]].rename(columns={CLUB: "录取社团"})  # 改名避免与报名表的「社团」同名列冲突
    m = src.merge(res, on=KEY, how="left")

    extra = set(res[KEY]) - set(src[KEY])
    if extra:
        print("[提示] 录取表中以下学生不在原始报名表内：" + "、".join(sorted(extra)))

    m[CLUB] = m.pop("录取社团").fillna(NO).replace("", NO)  # 录取 -> 社团名；未录取 -> 未被录取
    # 未录取排到最后，组内保持原顺序
    return m.assign(_no=m[CLUB].eq(NO)).sort_values("_no", kind="stable").drop(columns="_no")


def main():
    out = merge(read(SRC), read(RES))
    try:
        out.to_excel(OUT, index=False, sheet_name="录取结果")
    except PermissionError:
        raise SystemExit("无法写入 %s，请先关闭该文件（Excel / WPS 正打开着它）。" % OUT)

    # 美化：表头加粗填色、冻结首行、居中、自适应列宽，未录取标红
    wb = load_workbook(OUT)
    ws = wb.active
    ws.freeze_panes = "A2"
    for col in ws.columns:
        width = max(sum(2 if ord(ch) > 127 else 1 for ch in str(c.value or "")) for c in col) + 4
        ws.column_dimensions[col[0].column_letter].width = max(10, width)
        for c in col:
            c.alignment = Alignment(horizontal="center", vertical="center")
            if c.row == 1:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="4F81BD")
            elif c.value == NO:
                c.font = Font(color="C00000")
    wb.save(OUT)

    admitted = (out[CLUB] != NO).sum()
    print("合计 %d 人：录取 %d 人，未录取 %d 人\n已保存：%s" % (len(out), admitted, len(out) - admitted, OUT))


if __name__ == "__main__":
    main()
